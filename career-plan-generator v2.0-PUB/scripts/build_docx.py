#!/usr/bin/env python3
"""
career-plan-generator · docx 备份生成器（Python 版）

用法:
    python build_docx.py plan_data.json output.docx

依赖:
    pip install python-docx

设计目的:
    保证「在任何 Agent / 任何模型环境下都能生成 .docx 文件」。
    优先顺序：node + docx-js（首选）→ Python + python-docx（本脚本）→ Pandoc → 最后才退化为 md。

数据结构（与 build_docx.js 完全兼容）:
    {
      "studentName": "称呼",
      "subtitle": "—— 子标题 ——",
      "epigraph": "可选引言",
      "generatedAt": "2026-05",
      "planRange": "2026 — 2030",
      "portrait": "学生画像速写",
      "chapters": [
        {
          "title": "第X章 ...",
          "blocks": [
            {"type": "p", "text": "段落..."},
            {"type": "h2", "text": "二级标题"},
            {"type": "h3", "text": "三级标题"},
            {"type": "bullet", "items": ["...", "..."]},
            {"type": "numbered", "items": ["...", "..."]},
            {"type": "quote", "text": "引用块"},
            {"type": "table", "header": ["列1", "列2"], "rows": [["a","b"]]},
            {"type": "pagebreak"},
            {"type": "blank"}
          ]
        }
      ],
      "closingLetter": "写给你的一封信"
    }
"""

import json
import sys
from datetime import date

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("[build_docx.py] 缺少依赖：python-docx。请先运行 `pip install python-docx`。", file=sys.stderr)
    sys.exit(2)


# ---------- 配置 ----------
FONT_CN = "Microsoft YaHei"
FONT_EN = "Calibri"

COLOR = {
    "H1": RGBColor(0x1F, 0x4E, 0x79),
    "H2": RGBColor(0x2E, 0x75, 0xB6),
    "H3": RGBColor(0x40, 0x40, 0x40),
    "text": RGBColor(0x26, 0x26, 0x26),
    "mute": RGBColor(0x8C, 0x8C, 0x8C),
    "rule": RGBColor(0xB4, 0xC7, 0xE7),
    "table_head": RGBColor(0xD9, 0xE2, 0xF3),
    "quote_bg": "F2F7FB",
    "quote_bar": RGBColor(0x2E, 0x75, 0xB6),
}


# ---------- 辅助 ----------
def set_run_font(run, font=FONT_CN, size=11, bold=False, italic=False, color=None):
    """对一个 run 设置字体（含中文东亚字体）"""
    run.font.name = FONT_EN
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    # 中文字体设置
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), font)
    rFonts.set(qn("w:hAnsi"), FONT_EN)
    rFonts.set(qn("w:ascii"), FONT_EN)


def add_paragraph(doc, text, *, size=11, bold=False, italic=False,
                  color=None, align=None, space_before=0, space_after=4):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color or COLOR["text"])
    return p


def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    set_run_font(run, size=20, bold=True, color=COLOR["H1"])
    return p


def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=15, bold=True, color=COLOR["H2"])
    return p


def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=13, bold=True, color=COLOR["H3"])
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.4
    run = p.add_run(text)
    set_run_font(run, size=11, color=COLOR["text"])
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.4
    run = p.add_run(text)
    set_run_font(run, size=11, color=COLOR["text"])
    return p


def add_quote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    # 设置背景色 + 左边框
    pPr = p._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), COLOR["quote_bg"])
    pPr.append(shd)
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "2E75B6")
    pBdr.append(left)
    pPr.append(pBdr)
    run = p.add_run(text)
    set_run_font(run, size=11, italic=True, color=COLOR["H3"])
    return p


def add_table(doc, header, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 表头
    hdr = table.rows[0].cells
    for i, h in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(h))
        set_run_font(run, size=10, bold=True, color=COLOR["text"])
        # 表头背景
        tcPr = hdr[i]._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "D9E2F3")
        tcPr.append(shd)
    # 数据行
    for ri, row in enumerate(rows, start=1):
        cells = table.rows[ri].cells
        for ci, val in enumerate(row):
            cells[ci].text = ""
            p = cells[ci].paragraphs[0]
            run = p.add_run(str(val) if val is not None else "")
            set_run_font(run, size=10, color=COLOR["text"])
    return table


def add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)


# ---------- 块渲染分发 ----------
def render_block(doc, block):
    t = block.get("type", "")
    if t == "p":
        opts = block.get("opts", {}) or {}
        add_paragraph(
            doc, block.get("text", ""),
            size=opts.get("size", 11),
            bold=opts.get("bold", False),
            italic=opts.get("italics", False),
            align=opts.get("align"),
            color=COLOR.get(opts.get("color"), None) if isinstance(opts.get("color"), str) else None,
        )
    elif t == "h1":
        add_heading_1(doc, block.get("text", ""))
    elif t == "h2":
        add_heading_2(doc, block.get("text", ""))
    elif t == "h3":
        add_heading_3(doc, block.get("text", ""))
    elif t == "bullet":
        for item in block.get("items", []):
            add_bullet(doc, item)
    elif t == "numbered":
        for item in block.get("items", []):
            add_numbered(doc, item)
    elif t == "quote":
        add_quote(doc, block.get("text", ""))
    elif t == "table":
        add_table(doc, block.get("header", []), block.get("rows", []))
        add_paragraph(doc, " ")
    elif t == "pagebreak":
        add_page_break(doc)
    elif t == "blank":
        add_paragraph(doc, " ")
    else:
        add_paragraph(doc, f"[未知块类型：{t}]", color=RGBColor(0xC0, 0x00, 0x00))


# ---------- 主构造 ----------
def build(plan_data, out_path):
    doc = Document()

    # A4 + 边距
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # 页眉
    student_name = plan_data.get("studentName") or "你"
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run(f"{student_name}的职业规划成长计划书 · v1.4")
    set_run_font(hr, size=9, color=COLOR["mute"])

    # 页脚（页码）
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("— 第 ")
    set_run_font(fr, size=9, color=COLOR["mute"])
    # 页码字段
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    fp._p.append(field)
    fr2 = fp.add_run(" 页 —")
    set_run_font(fr2, size=9, color=COLOR["mute"])

    # ========== 封面 ==========
    add_paragraph(doc, "", space_before=80)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("职业规划成长计划书")
    set_run_font(run, size=28, bold=True, color=COLOR["H1"])

    subtitle = plan_data.get("subtitle") or "—— 定制版职业规划 ——"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(60)
    run = p.add_run(subtitle)
    set_run_font(run, size=14, color=COLOR["H2"])

    epigraph = plan_data.get("epigraph")
    if epigraph:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(60)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f"「{epigraph}」")
        set_run_font(run, size=12, italic=True, color=RGBColor(0x59, 0x59, 0x59))

    add_paragraph(doc, " ", space_before=60)
    add_paragraph(doc, f"署名：{student_name}", size=12, align="center", space_before=4, space_after=4)

    plan_range = plan_data.get("planRange")
    if plan_range:
        add_paragraph(doc, f"规划周期：{plan_range}", size=12, align="center", space_before=4, space_after=4)
    generated_at = plan_data.get("generatedAt")
    if generated_at:
        add_paragraph(doc, f"拟定日期：{generated_at}", size=12, align="center", space_before=4, space_after=4)

    # v1.4 封面底部国科声明
    add_paragraph(doc, " ", space_before=80)
    add_paragraph(
        doc, "本报告由国科职规体系 + AI 生成，受限于模型能力，仅供参看。",
        size=8, italic=True, align="center", color=COLOR["mute"], space_before=80, space_after=2
    )
    add_paragraph(
        doc, "如有需求，可向国科职规老师获取专业能力测评与职规服务。",
        size=8, italic=True, align="center", color=COLOR["mute"], space_before=2, space_after=4
    )

    add_page_break(doc)

    # ========== 学生画像速写 ==========
    portrait = plan_data.get("portrait")
    if portrait:
        add_heading_1(doc, "学生画像速写")
        add_quote(doc, portrait)
        add_page_break(doc)

    # ========== 各章 ==========
    for chapter in plan_data.get("chapters", []):
        add_heading_1(doc, chapter.get("title", ""))
        for block in chapter.get("blocks", []):
            render_block(doc, block)
        add_page_break(doc)

    # ========== 第十三章 · 5. 写给你的一封信 ==========
    closing = plan_data.get("closingLetter")
    if closing:
        add_heading_1(doc, "写给你的一封信")
        add_quote(doc, closing)
        add_paragraph(doc, " ")
        add_paragraph(
            doc, "—— 本计划书 v1.4 完 ——",
            italic=True, align="center", color=COLOR["mute"], space_before=12, space_after=4
        )

    doc.save(out_path)
    print(f"[build_docx.py] 已生成：{out_path}")


def main():
    if len(sys.argv) < 3:
        print("用法：python build_docx.py <plan_data.json> <output.docx>", file=sys.stderr)
        sys.exit(1)
    with open(sys.argv[1], "r", encoding="utf-8") as f:
        plan_data = json.load(f)
    build(plan_data, sys.argv[2])


if __name__ == "__main__":
    main()
